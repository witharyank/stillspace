import os
import re

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Error: python-docx not installed. Please run: pip install python-docx")
    exit(1)

def create_docx_from_md(md_file, docx_file):
    if not os.path.exists(md_file):
        print(f"Error: {md_file} not found.")
        return

    doc = Document()
    
    # Setup styling globally
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    in_code_block = False
    
    print(f"Parsing {md_file} and generating DOCX...")

    for line in lines:
        stripped = line.strip()
        
        if stripped.startswith("```"):
            in_code_block = not in_code_block
            continue
            
        if in_code_block:
            p = doc.add_paragraph(line)
            p.runs[0].font.name = 'Courier New'
            p.runs[0].font.size = Pt(10)
            continue
            
        if stripped == "---":
            doc.add_page_break()
            continue
            
        if stripped.startswith("#"):
            level = len(line) - len(line.lstrip('#'))
            # Force level up to 3 max for python-docx built in styles easily
            level = min(level, 3) 
            header_text = line.replace('#', '').strip()
            
            h = doc.add_heading(header_text, level=level)
            h.runs[0].font.name = 'Arial'
            if level == 1:
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                h.runs[0].bold = True
            continue
            
        if stripped.startswith("- "):
            p = doc.add_paragraph(stripped[1:].strip(), style='List Bullet')
        elif re.match(r'^\d+\.', stripped):
            p = doc.add_paragraph(stripped, style='List Number')
        elif len(stripped) > 0:
            p = doc.add_paragraph()
            # Simple bold formatting parser
            parts = re.split(r'(\*\*.*?\*\*)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part.replace('**', ''))
                    run.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    run = p.add_run(part.replace('*', ''))
                    run.italic = True
                else:
                    # Clean up standard generic math / literal symbols for raw word compatibility safely
                    clean_part = part.replace('$', '').replace('`', '')
                    p.add_run(clean_part)
        else:
            doc.add_paragraph()

    doc.save(docx_file)
    print(f"Success! Document saved successfully as: {os.path.abspath(docx_file)}")

if __name__ == "__main__":
    create_docx_from_md('report.md', 'StillSpace_DAA_Project_Report.docx')
